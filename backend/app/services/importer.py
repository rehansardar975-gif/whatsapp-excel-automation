"""Excel/CSV import: read -> detect columns -> normalize -> deduplicate -> check suppression -> summarize."""
import csv
import io
import re
from datetime import datetime, timezone

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from .. import config
from ..models import Contact, ImportBatch, Suppression
from . import activity
from .phone import normalize_phone


class ImportErrorForUser(ValueError):
    """An error message that can be shown to a non-technical user as-is."""


FIELDS = ["full_name", "first_name", "last_name", "company", "phone", "country", "email", "city", "industry", "source", "opt_in"]
FIELD_LABELS = {
    "full_name": "Full name", "first_name": "First name", "last_name": "Last name", "company": "Business",
    "phone": "Phone", "country": "Country", "email": "Email", "city": "City", "industry": "Industry",
    "source": "Source", "opt_in": "WhatsApp opt-in",
}
SYNONYMS = {
    "phone": ["phone", "phone number", "mobile", "mobile number", "whatsapp", "whatsapp number", "cell", "cell phone",
              "tel", "telephone", "contact number", "number", "phone no", "mobile no", "msisdn"],
    "first_name": ["first name", "firstname", "first", "given name", "fname"],
    "last_name": ["last name", "lastname", "surname", "family name", "lname"],
    "full_name": ["name", "full name", "contact name", "contact", "customer name", "client name", "person"],
    "company": ["company", "business", "business name", "company name", "organization", "organisation", "shop", "store", "brand", "firm"],
    "country": ["country", "country name", "nation"],
    "email": ["email", "e-mail", "email address", "mail"],
    "city": ["city", "town", "location", "area"],
    "industry": ["industry", "category", "sector", "business type", "type"],
    "source": ["source", "lead source", "origin", "channel"],
    "opt_in": ["opt in", "opt-in", "optin", "consent", "whatsapp opt in", "whatsapp consent", "subscribed", "opted in", "marketing consent", "agreed"],
}
YES = {"yes", "y", "true", "1", "opted in", "opt-in", "opt in", "subscribed", "agreed", "x", "✓", "consented"}
NO = {"no", "n", "false", "0", "opted out", "unsubscribed", "stop", "do not contact", "dnc"}


def _norm_header(h) -> str:
    return re.sub(r"[\s_\-.:#]+", " ", str(h or "")).strip().lower()


# ---------- reading ----------
def read_table(file_name: str, content: bytes) -> tuple[list[str], list[list]]:
    if not content:
        raise ImportErrorForUser("The file is empty.")
    if len(content) > config.MAX_UPLOAD_BYTES:
        raise ImportErrorForUser(f"The file is larger than {config.MAX_UPLOAD_BYTES // (1024 * 1024)} MB.")
    name = file_name.lower()
    if name.endswith((".xlsx", ".xlsm")):
        grid = _read_xlsx(content)
    elif name.endswith((".csv", ".txt")):
        grid = _read_csv(content)
    elif name.endswith(".xls"):
        raise ImportErrorForUser("Old Excel format (.xls) is not supported. In Excel, choose File → Save As → .xlsx and upload again.")
    else:
        raise ImportErrorForUser("Please upload an Excel (.xlsx) or CSV file.")
    grid = [r for r in grid if any(str(c).strip() for c in r if c is not None)]
    if not grid:
        raise ImportErrorForUser("No rows were found in the file.")
    h_idx = _find_header_row(grid)
    headers = [str(c).strip() if c is not None and str(c).strip() else f"Column {i + 1}" for i, c in enumerate(grid[h_idx])]
    rows = [list(r) + [None] * (len(headers) - len(r)) for r in grid[h_idx + 1:]]
    rows = [[_jsonable(c) for c in r[: len(headers)]] for r in rows]
    if not rows:
        raise ImportErrorForUser("The file has a header row but no contacts.")
    if len(rows) > config.MAX_UPLOAD_ROWS:
        raise ImportErrorForUser(f"The file has {len(rows):,} rows. The limit is {config.MAX_UPLOAD_ROWS:,} per upload; please split it.")
    return headers, rows


def _jsonable(c):
    if isinstance(c, datetime):
        return c.date().isoformat()
    if isinstance(c, (int, float, str, bool)) or c is None:
        return c
    return str(c)


def _read_xlsx(content: bytes) -> list[list]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise ImportErrorForUser("This Excel file could not be opened. It may be damaged or password-protected.")
    # first sheet that has data
    for ws in wb.worksheets:
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
        if any(any(c is not None and str(c).strip() for c in r) for r in grid):
            wb.close()
            return grid
    wb.close()
    return []


def _read_csv(content: bytes) -> list[list]:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = content.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        delim = dialect.delimiter
    except csv.Error:
        delim = ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delim)]


def _find_header_row(grid: list[list]) -> int:
    """Files often have a title row or two above the real header. Pick the row that looks most like headers."""
    best, best_score = 0, 0
    for i, row in enumerate(grid[:10]):
        score = sum(1 for c in row if _match_field(_norm_header(c)))
        if score > best_score:
            best, best_score = i, score
    return best


def _match_field(h: str) -> str | None:
    if not h:
        return None
    for field, words in SYNONYMS.items():
        if h in words:
            return field
    for field, words in SYNONYMS.items():
        if any(len(w) > 3 and w in h for w in words):
            return field
    return None


def detect_mapping(headers: list[str]) -> dict[str, str | None]:
    """Return {field: header or None}. Each header is used at most once."""
    mapping: dict[str, str | None] = {f: None for f in FIELDS}
    for h in headers:
        f = _match_field(_norm_header(h))
        if f and mapping[f] is None:
            mapping[f] = h
    # a single "Name" column with explicit first/last columns present is ambiguous: prefer first/last
    if mapping["first_name"] and mapping["full_name"]:
        mapping["full_name"] = None
    return mapping


# ---------- analysis ----------
def _cell(row: list, headers: list[str], header: str | None):
    if not header or header not in headers:
        return None
    return row[headers.index(header)]


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def _title(s: str) -> str:
    return s.title() if s and (s.isupper() or s.islower()) else s


def parse_opt_in(v) -> bool | None:
    t = _text(v).lower()
    if t in YES:
        return True
    if t in NO:
        return False
    return None


def analyze(db: Session, headers: list[str], rows: list[list], mapping: dict, default_country: str,
            consent_confirmed: bool = False) -> tuple[list[dict], dict]:
    if not mapping.get("phone"):
        raise ImportErrorForUser("We couldn't find a phone number column. Choose which column has the phone numbers.")
    existing = set(db.scalars(select(Contact.phone).where(Contact.phone.is_not(None))))
    suppressed = {s.phone: s.reason for s in db.scalars(select(Suppression))}
    seen: dict[str, int] = {}
    out = []
    for i, row in enumerate(rows, start=1):
        g = lambda f: _cell(row, headers, mapping.get(f))  # noqa: E731
        first, last = _title(_text(g("first_name"))), _title(_text(g("last_name")))
        full = _title(_text(g("full_name")))
        if full and not first:
            parts = full.split(" ", 1)
            first, last = parts[0], (parts[1] if len(parts) > 1 else last)
        country_raw = _text(g("country"))
        ph = normalize_phone(g("phone"), default_country, country_raw)
        opt = parse_opt_in(g("opt_in"))
        rec = {
            "row": i, "first_name": first, "last_name": last, "company": _text(g("company")),
            "phone_raw": _text(g("phone")), "phone": ph.e164, "country": ph.region or "",
            "email": _text(g("email")).lower(), "city": _title(_text(g("city"))),
            "industry": _text(g("industry")), "source": _text(g("source")),
            "opted_in": opt is True or (opt is None and consent_confirmed),
            "opt_in_value": opt, "status": "ready", "reason": "",
        }
        if not (first or last or rec["company"] or rec["phone_raw"]):
            continue  # fully blank row
        if ph.status == "missing":
            rec.update(status="missing", reason="No phone number")
        elif ph.status == "invalid":
            rec.update(status="invalid", reason=f"{ph.reason}: {rec['phone_raw']}")
        elif ph.e164 in seen:
            rec.update(status="duplicate", reason=f"Same number as row {seen[ph.e164]}")
        elif ph.e164 in suppressed:
            rec.update(status="do_not_contact", reason=f"On Do Not Contact list ({suppressed[ph.e164]})")
        elif opt is False:
            rec.update(status="do_not_contact", reason="Marked as opted out in the file")
        elif ph.e164 in existing:
            rec.update(status="existing", reason="Already in your contacts")
        elif not rec["opted_in"]:
            rec.update(status="needs_opt_in", reason="No WhatsApp opt-in on record")
        if ph.e164 and ph.e164 not in seen:
            seen[ph.e164] = i
        out.append(rec)
    summary = {k: 0 for k in ["total", "ready", "needs_opt_in", "duplicate", "existing", "invalid", "missing", "do_not_contact"]}
    summary["total"] = len(out)
    for r in out:
        summary[r["status"]] += 1
    summary["will_import"] = summary["ready"] + summary["needs_opt_in"]
    summary["has_opt_in_column"] = bool(mapping.get("opt_in"))
    return out, summary


# ---------- import ----------
def commit(db: Session, batch: ImportBatch, consent_confirmed: bool, via: str = "app") -> dict:
    records, summary = analyze(db, batch.headers, batch.rows, batch.mapping, batch.default_country, consent_confirmed)
    now = datetime.now(timezone.utc)
    created = 0
    for r in records:
        if r["status"] not in ("ready", "needs_opt_in"):
            if r["status"] == "do_not_contact" and r["opt_in_value"] is False and r["phone"]:
                from .suppression import add_suppression
                add_suppression(db, r["phone"], "Opted out (from imported file)", "import", log_event=False)
            continue
        consent_src = ""
        if r["opted_in"]:
            consent_src = "Opt-in column in file" if r["opt_in_value"] is True else "Confirmed by user at import"
        db.add(Contact(
            first_name=r["first_name"], last_name=r["last_name"], company=r["company"], phone_raw=r["phone_raw"],
            phone=r["phone"], country=r["country"], email=r["email"], city=r["city"], industry=r["industry"],
            source=r["source"] or ("Contact collection" if batch.source == "collection" else "Excel import"),
            phone_status="valid", opted_in=r["opted_in"], consent_source=consent_src,
            consent_at=now if r["opted_in"] else None, import_batch_id=batch.id,
            last_activity=f"Imported from {batch.file_name}", last_activity_at=now,
        ))
        created += 1
    batch.status = "imported"
    batch.imported_at = now
    summary["imported"] = created
    batch.summary = summary
    batch.rows = []  # raw file data is no longer needed once imported
    activity.log(db, "import", f"Imported {created} contacts from {batch.file_name}" + (f" via {via}" if via != "app" else ""), "success",
                 batch_id=batch.id, **{k: summary[k] for k in ("total", "duplicate", "invalid", "missing", "do_not_contact", "existing")})
    if summary["duplicate"]:
        activity.log(db, "duplicate", f"{summary['duplicate']} duplicate rows removed from {batch.file_name}", "warning", batch_id=batch.id)
    if summary["invalid"] or summary["missing"]:
        activity.log(db, "invalid", f"{summary['invalid']} invalid and {summary['missing']} missing phone numbers in {batch.file_name}", "warning", batch_id=batch.id)
    db.commit()
    return summary
