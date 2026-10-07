"""End-to-end API workflow: upload -> review -> import -> message -> campaign -> process -> results."""
import csv
import io

from openpyxl import Workbook, load_workbook

from app.providers.demo import _bucket


def xlsx_bytes(rows):
    wb = Workbook()
    ws = wb.active
    ws.append(["Customer Contacts — exported"])  # title row above the header
    ws.append([])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


HEADER = ["Full Name", "Company Name", "WhatsApp Number", "Country", "Opt-in"]
ROWS = [
    HEADER,
    ["sara khan", "Palm Grill", "050 111 2233", "UAE", "Yes"],
    ["Omar Haddad", "Skyline Realty", "+971 55 444 5566", "UAE", "yes"],
    ["Sara Khan", "Palm Grill", "00971501112233", "UAE", "Yes"],       # duplicate of row 1
    ["Bad Number", "Test LLC", "12345", "UAE", "Yes"],                 # invalid
    ["No Phone", "Test LLC", "", "UAE", "Yes"],                        # missing
    ["Opted Out", "Cedar Café", "0551231234", "Saudi Arabia", "No"],   # opted out in file
    ["No Consent", "Glow Salon", "0509998877", "UAE", ""],             # needs opt-in
]


def upload(c, name, content, country="AE"):
    return c.post("/api/imports", files={"file": (name, content)}, data={"default_country": country})


def test_excel_upload_detects_columns_and_classifies(fresh):
    r = upload(fresh, "leads.xlsx", xlsx_bytes(ROWS))
    assert r.status_code == 200, r.text
    d = r.json()
    m = d["mapping"]
    assert (m["full_name"], m["company"], m["phone"], m["country"], m["opt_in"]) == \
           ("Full Name", "Company Name", "WhatsApp Number", "Country", "Opt-in")
    s = d["summary"]
    assert (s["total"], s["ready"], s["duplicate"], s["invalid"], s["missing"], s["do_not_contact"], s["needs_opt_in"]) == (7, 2, 1, 1, 1, 1, 1)
    first = d["records"][0]
    assert first["first_name"] == "Sara" and first["last_name"] == "Khan" and first["phone"] == "+971501112233"
    assert d["records"][5]["phone"] == "+966551231234"  # country column used


def test_csv_upload_semicolon_and_mapping_change(fresh):
    text = "Name;Business;Tel;Mobile\nAli Raza;Raza Traders;;0501234567\n"
    d = upload(fresh, "list.csv", text.encode("utf-8")).json()
    assert d["mapping"]["phone"] == "Tel"
    assert d["summary"]["missing"] == 1
    r = fresh.put(f"/api/imports/{d['id']}", json={"mapping": {**d["mapping"], "phone": "Mobile"}})
    assert r.json()["summary"]["needs_opt_in"] == 1  # valid number now, but no opt-in column
    r = fresh.put(f"/api/imports/{d['id']}", json={"mapping": {**d["mapping"], "phone": "Mobile"}, "consent_confirmed": True})
    assert r.json()["summary"]["ready"] == 1
    assert fresh.put(f"/api/imports/{d['id']}", json={"mapping": {"phone": "Nope"}}).status_code == 422


def test_upload_errors_are_friendly(fresh):
    assert "empty" in upload(fresh, "a.csv", b"").json()["detail"]
    assert ".xlsx" in upload(fresh, "a.xls", b"xx").json()["detail"]
    assert "Excel (.xlsx) or CSV" in upload(fresh, "a.pdf", b"xx").json()["detail"]
    assert upload(fresh, "broken.xlsx", b"not a zip").status_code == 422
    d = upload(fresh, "nophone.csv", b"Name,City\nA,B\n").json()
    assert "phone number column" in d["summary"]["error"]


def test_import_campaign_suppression_end_to_end(fresh):
    c = fresh
    before = c.get("/api/contacts").json()["total"]
    d = upload(c, "leads.xlsx", xlsx_bytes(ROWS)).json()
    res = c.post(f"/api/imports/{d['id']}/commit", json={"consent_confirmed": False}).json()
    assert res["imported"] == 3  # 2 ready + 1 needs opt-in
    assert c.post(f"/api/imports/{d['id']}/commit", json={}).status_code == 409  # no double import
    # the opted-out number went to the Do Not Contact list
    assert any(s["phone"] == "+966551231234" for s in c.get("/api/suppressions").json())

    # re-uploading the same file: everything already known
    s2 = upload(c, "again.xlsx", xlsx_bytes(ROWS)).json()["summary"]
    assert s2["existing"] == 3 and s2["ready"] == 0 and s2["do_not_contact"] == 1

    contacts = c.get("/api/contacts", params={"batch_id": d["id"]}).json()
    assert contacts["total"] == 3
    assert c.get("/api/contacts").json()["total"] == before + 3

    # search / filter / sort
    assert c.get("/api/contacts", params={"q": "skyline", "batch_id": d["id"]}).json()["total"] == 1
    st = c.get("/api/contacts", params={"batch_id": d["id"], "status": "needs_opt_in"}).json()
    assert st["total"] == 1 and st["items"][0]["reason"] == "No WhatsApp opt-in on record"
    names = [x["company"] for x in c.get("/api/contacts", params={"batch_id": d["id"], "sort": "company", "order": "asc"}).json()["items"]]
    assert names == sorted(names)

    # message + campaign
    t = c.post("/api/templates", json={"name": "Test offer", "body": "Hi {{first_name}} from {{company}}"}).json()
    prev = c.post("/api/campaigns/audience-preview", json={"batch_id": d["id"]}).json()
    assert prev["ready"] == 2 and prev["excluded"] == 1
    camp = c.post("/api/campaigns", json={"name": "Test", "template_id": t["id"], "audience": {"batch_id": d["id"]}}).json()
    assert camp["status"] == "draft" and camp["report"]["queued"] == 2 and camp["report"]["skipped"] == 1

    # a contact opts out AFTER the campaign was created -> must be skipped at send time
    omar = next(x for x in contacts["items"] if x["company"] == "Skyline Realty")
    assert c.post("/api/suppressions", json={"phone": omar["phone"], "reason": "Asked by phone"}).status_code == 201

    from app.services.campaigns import launch
    c.post(f"/api/campaigns/{camp['id']}/start")
    launch(camp["id"], sync=True)  # wait for the background queue deterministically
    rep = c.get(f"/api/campaigns/{camp['id']}").json()
    assert rep["status"] == "completed"
    r = rep["report"]
    assert r["remaining"] == 0 and r["skipped"] == 2 and r["processed"] == 3
    msgs = c.get(f"/api/campaigns/{camp['id']}/messages").json()["items"]
    by_name = {m["contact_name"]: m for m in msgs}
    assert by_name["Omar Haddad"]["status"] == "skipped" and "Do Not Contact" in by_name["Omar Haddad"]["reason"]
    assert by_name["Sara Khan"]["body"] == "Hi Sara from Palm Grill"
    assert c.post(f"/api/campaigns/{camp['id']}/start").status_code == 409

    # future campaigns: suppressed contact is never queued
    camp2 = c.post("/api/campaigns", json={"name": "Again", "template_id": t["id"], "audience": {"batch_id": d["id"]}}).json()
    assert camp2["report"]["queued"] == 1

    # opt-in cannot be re-added while on Do Not Contact
    assert c.post(f"/api/contacts/{omar['id']}/consent", json={"opted_in": True}).status_code == 409

    # export
    x = c.get(f"/api/campaigns/{camp['id']}/export?format=xlsx")
    ws = load_workbook(io.BytesIO(x.content)).active
    assert ws.max_row == 4 and ws["A1"].value == "Contact"
    rows = list(csv.reader(io.StringIO(c.get("/api/contacts/export?format=csv").content.decode("utf-8-sig"))))
    assert rows[0][0] == "First name" and len(rows) > 3


def test_demo_outcomes_success_failure_retry_optout(fresh):
    """Seeded campaign exercises every outcome; check each one is classified correctly."""
    c = fresh
    camp = c.get("/api/campaigns").json()[0]
    r = camp["report"]
    assert r["sent"] > 0 and r["delivered"] > 0 and r["failed"] > 0 and r["skipped"] > 0 and r["opted_out"] > 0 and r["retried"] > 0
    assert r["sent"] + r["failed"] + r["skipped"] == r["total"]
    msgs = c.get(f"/api/campaigns/{camp['id']}/messages", params={"page_size": 200}).json()["items"]
    for m in msgs:
        b = _bucket(m["phone"]) if m["phone"].startswith("+") else None
        if m["status"] == "skipped":
            continue
        if b is not None and b < 3:
            assert m["status"] == "failed" and m["error_code"] == "131026"
        elif b == 3:
            assert m["status"] == "failed" and m["opted_out"]
        elif b == 4:
            assert m["status"] == "failed" and m["retry_count"] == 2
        elif b in (5, 6, 7):
            assert m["status"] == "delivered" and m["retry_count"] == 1
        elif b in (8, 9):
            assert m["status"] == "delivered" and m["opted_out"]
        elif b in (10, 11):
            assert m["status"] == "sent"
        else:
            assert m["status"] == "delivered"
    # every opted-out recipient is on the Do Not Contact list
    sup = {s["phone"] for s in c.get("/api/suppressions").json()}
    assert all(m["phone"] in sup for m in msgs if m["opted_out"])


def test_campaign_validation(fresh):
    c = fresh
    assert c.post("/api/campaigns", json={"name": "X", "template_id": 9999}).status_code == 422
    t = c.get("/api/templates").json()[0]
    assert c.post("/api/campaigns", json={"name": "X", "template_id": t["id"], "audience": {"country": "ZZ"}}).status_code == 422
    assert c.post("/api/templates", json={"name": "Bad", "body": "Hi {{nickname}}"}).status_code == 422
    assert c.get("/api/campaigns/9999").status_code == 404


def test_webhook_verification_status_and_stop(fresh):
    c = fresh
    assert c.get("/api/webhooks/whatsapp", params={"hub.mode": "subscribe", "hub.verify_token": "demo-verify-token", "hub.challenge": "42"}).text == "42"
    assert c.get("/api/webhooks/whatsapp", params={"hub.mode": "subscribe", "hub.verify_token": "wrong", "hub.challenge": "42"}).status_code == 403
    camp = c.get("/api/campaigns").json()[0]
    msg = next(m for m in c.get(f"/api/campaigns/{camp['id']}/messages", params={"status": "delivered", "page_size": 200}).json()["items"] if not m["opted_out"])
    # an out-of-order "sent" must not downgrade a delivered message
    payload = {"entry": [{"changes": [{"value": {"statuses": [{"id": msg["provider_message_id"], "status": "sent"}]}}]}]}
    assert c.post("/api/webhooks/whatsapp", json=payload).json()["applied"] == 1
    m2 = next(m for m in c.get(f"/api/campaigns/{camp['id']}/messages", params={"page_size": 200}).json()["items"] if m["id"] == msg["id"])
    assert m2["status"] == "delivered"
    # inbound STOP reply -> Do Not Contact
    stop = {"entry": [{"changes": [{"value": {"messages": [{"from": msg["phone"].lstrip("+"), "type": "text", "text": {"body": "STOP"},
                                                             "context": {"id": msg["provider_message_id"]}}]}}]}]}
    c.post("/api/webhooks/whatsapp", json=stop)
    assert msg["phone"] in {s["phone"] for s in c.get("/api/suppressions").json()}


def test_collection_clean_export_and_add_as_needs_opt_in(fresh):
    c = fresh
    run = c.post("/api/collection/runs", json={"industry": "Restaurants", "city": "Dubai", "limit": 40}).json()
    s = run["summary"]
    assert s["total"] > 40 and s["duplicate"] >= 1 and s["will_import"] > 0
    assert [st["label"] for st in run["steps"]] == ["Collected", "Valid phone", "Unique", "New to your contacts"]
    x = c.get(f"/api/collection/runs/{run['id']}/export?format=xlsx")
    assert x.status_code == 200 and load_workbook(io.BytesIO(x.content)).active["I1"].value == "Validation"
    res = c.post(f"/api/collection/runs/{run['id']}/add").json()
    added = c.get("/api/contacts", params={"batch_id": run["id"]}).json()
    assert res["imported"] == added["total"] and all(i["status"] == "needs_opt_in" for i in added["items"])
    assert c.post("/api/collection/runs", json={"industry": "Hacking", "city": "Dubai"}).status_code == 422


def test_dashboard_and_activity(fresh):
    d = fresh.get("/api/dashboard").json()
    assert d["mode"] == "demo" and d["contacts"] > 0 and d["campaigns"] == 1 and d["opt_outs"] >= 1
    a = fresh.get("/api/activity", params={"type": "import"}).json()
    assert a["total"] >= 1 and all(i["type"] == "import" for i in a["items"])


def test_timestamps_carry_utc_offset(fresh):
    """Browsers in other time zones must not read UTC times as local time."""
    c = fresh.get("/api/campaigns").json()[0]
    assert c["created_at"].endswith("+00:00") and c["completed_at"].endswith("+00:00")
    a = fresh.get("/api/activity").json()["items"][0]
    assert a["created_at"].endswith("+00:00")
